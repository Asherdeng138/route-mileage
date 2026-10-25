"""
独立核对导出的 Word / Excel（不使用 App 自己的代码）：
  - openpyxl / python-docx 能否正常打开
  - Excel：工作表、表头、公式、SUM 范围、打印区域、数字格式
  - Word：页面设置、标题段落、行程段落、图片
  - Word 与 Excel 的日期、行程数、里程、金额是否一致

用法：python -I scripts/verify-exports.py samples/SAMPLE-Mileage-Claim-2026-10.docx samples/SAMPLE-Mileage-Details-2026-10.xlsx
依赖：pip install openpyxl python-docx
"""
import re
import sys
from decimal import ROUND_HALF_UP, Decimal

import docx
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
docx_path, xlsx_path = sys.argv[1], sys.argv[2]
problems = []


def check(cond, msg):
    if not cond:
        problems.append(msg)


# ---------- Excel ----------
wb = openpyxl.load_workbook(xlsx_path)  # 公式
wbv = openpyxl.load_workbook(xlsx_path, data_only=True)  # 缓存值
check(wb.sheetnames == ['Sheet1'], f'sheet names {wb.sheetnames}')
ws, wsv = wb['Sheet1'], wbv['Sheet1']
check(re.fullmatch(r'\d{4} [A-Z][a-z]+ mileage', str(ws['A1'].value)), f'A1 title {ws["A1"].value!r}')
check([ws['A2'].value, ws['B2'].value, ws['C2'].value] == ['Date', 'Miles', 'Cost'], 'header row')
check(ws['A1'].font.b and ws['A2'].font.b, 'title / header bold')
check(ws['C3'].number_format == '"£"#,##0.00', f'cost number format {ws["C3"].number_format!r}')

rows = []
r = 3
while ws.cell(r, 1).value and re.match(r'\d{4}/\d{2}/\d{2} ', str(ws.cell(r, 1).value)):
    a, b, c = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
    check(isinstance(c, str) and c.startswith(f'=') and f'B{r}' in c, f'C{r} formula {c!r}')
    check(ws.cell(r, 1).alignment.wrap_text, f'A{r} wrap text')
    rows.append((a, Decimal(str(b)), Decimal(str(wsv.cell(r, 3).value))))
    r += 1
last = r - 1
labels = {str(ws.cell(i, 1).value): i for i in range(1, ws.max_row + 1) if ws.cell(i, 1).value}
month_row = next(i for k, i in labels.items() if k.startswith('Miles in ') and i > last)
check(ws.cell(month_row, 2).value == f'=SUM(B3:B{last})', f'month miles formula {ws.cell(month_row, 2).value}')
check(ws.cell(month_row, 3).value == f'=SUM(C3:C{last})', f'month cost formula {ws.cell(month_row, 3).value}')
year_row = labels.get('Takes me to miles this year')
check(year_row is not None, 'year total row')
check(ws.print_area and f'$C${year_row}' in str(ws.print_area), f'print area {ws.print_area}')
check(ws.print_title_rows == '$1:$2', f'print titles {ws.print_title_rows}')
check(ws.page_setup.paperSize in ('9', 9) and ws.page_setup.orientation == 'portrait', 'A4 portrait')

x_miles = sum(m for _, m, _ in rows)
x_cost = sum(c for _, _, c in rows).quantize(Decimal('0.01'), ROUND_HALF_UP)
check(Decimal(str(wsv.cell(month_row, 2).value)) == x_miles, 'cached month miles = sum of rows')

# ---------- Word ----------
d = docx.Document(docx_path)
sec = d.sections[0]
check((sec.page_width, sec.page_height) == (7560310, 10692130), f'page size {sec.page_width}x{sec.page_height}')
check(all(m == 914400 for m in (sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin)), 'margins 2.54 cm')
paras = [p.text for p in d.paragraphs]
check(paras[0].startswith('Mileage for '), f'title {paras[0]!r}')
check(paras[1].startswith('First 10,000 miles paid at £'), f'rate line {paras[1]!r}')
m_total = re.fullmatch(r'Miles total ([\d.]+) miles', paras[2])
m_amount = re.fullmatch(r'Amount to claim £([\d,]+\.\d\d)', paras[3])
m_year = re.fullmatch(r'Takes me to ([\d.]+) miles this year', paras[4])
check(m_total and m_amount and m_year, f'summary paragraphs {paras[2:5]}')
check(all(r.bold for r in d.paragraphs[0].runs), 'title bold')

word_trips = []
for p in paras[5:]:
    m = re.fullmatch(r'((?:\d{4}/\d{2}/\d{2})(?:&\d{4}/\d{2}/\d{2})*) (.+) ([\d.]+) miles( each)?', p)
    if m:
        for date in m.group(1).split('&'):
            word_trips.append((f'{date} {m.group(2)}', Decimal(m.group(3))))
images = [s for s in d.inline_shapes]
check(len(images) > 0, 'images present')
for s in images:
    check(s.width <= 5731510 and s.height <= 5940000, f'image size {s.width}x{s.height}')

# ---------- 一致性 ----------
check(len(word_trips) == len(rows), f'trip count word {len(word_trips)} vs excel {len(rows)}')
check(sorted(word_trips) == sorted((a, b) for a, b, _ in rows), 'trip dates / descriptions / miles differ')
if m_total:
    check(Decimal(m_total.group(1)) == x_miles, f'miles word {m_total.group(1)} vs excel {x_miles}')
if m_amount:
    check(Decimal(m_amount.group(1).replace(',', '')) == x_cost, f'amount word {m_amount.group(1)} vs excel {x_cost}')
if m_year and year_row:
    check(Decimal(m_year.group(1)) == Decimal(str(wsv.cell(year_row, 2).value)), 'year total differs')

print(f'Excel: {len(rows)} trips, {x_miles} miles, £{x_cost}; Word: {len(word_trips)} trips, {len(images)} images, {len(paras)} paragraphs')
if problems:
    print('PROBLEMS:')
    for p in problems:
        print(' -', p)
    sys.exit(1)
print('OK: Word and Excel are structurally valid and consistent')
